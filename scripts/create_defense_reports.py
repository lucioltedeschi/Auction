from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
TECH_OUT = DOCS / "Informe_Tecnico_AuctIO_Entrega3.docx"
SUMMARY_OUT = DOCS / "Resumen_Cambios_Subastas_Defensa.docx"
NAVY, GOLD, CREAM, PALE, MUTED, GREEN, RED = "071827", "A8872F", "F3F0E8", "E8EEF5", "475569", "166534", "991B1B"

def font(run, size=10.5, color=NAVY, bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size, run.font.color.rgb = Pt(size), RGBColor.from_string(color)
    run.bold, run.italic = bold, italic

def shade(cell, fill):
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); cell._tc.get_or_add_tcPr().append(shd)

def geometry(table, widths):
    table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW")) or OxmlElement("w:tblW")
    if tbl_w.getparent() is None: tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths))); tbl_w.set(qn("w:type"), "dxa")
    ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), "120"); ind.set(qn("w:type"), "dxa"); tbl_pr.append(ind)
    grid = table._tbl.tblGrid
    for node in list(grid): grid.remove(node)
    for width in widths:
        node = OxmlElement("w:gridCol"); node.set(qn("w:w"), str(width)); grid.append(node)
    for row in table.rows:
        cant = OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(cant)
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_w = OxmlElement("w:tcW"); tc_w.set(qn("w:w"), str(widths[index])); tc_w.set(qn("w:type"), "dxa"); cell._tc.get_or_add_tcPr().append(tc_w)
            margins = OxmlElement("w:tcMar")
            for edge, value in (("top",80),("bottom",80),("start",120),("end",120)):
                n = OxmlElement("w:" + edge); n.set(qn("w:w"), str(value)); n.set(qn("w:type"), "dxa"); margins.append(n)
            cell._tc.get_or_add_tcPr().append(margins)

def table(doc, headers, rows, widths, font_size=8.7):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for i, value in enumerate(headers):
        shade(t.rows[0].cells[i], PALE)
        p = t.rows[0].cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        font(p.add_run(value), 9, NAVY, True)
    tr_header = OxmlElement("w:tblHeader"); t.rows[0]._tr.get_or_add_trPr().append(tr_header)
    for r_index, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            if r_index % 2: shade(cells[i], "F8FAFC")
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            font(p.add_run(str(value)), font_size)
    geometry(t, widths)
    return t

def heading(doc, text, level=1, page=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True; p.paragraph_format.page_break_before = page
    font(p.add_run(text), 16 if level == 1 else 13 if level == 2 else 12, GOLD if level == 1 else NAVY, True)
    return p

def paragraph(doc, text, size=10.5, color=NAVY, bold=False, italic=False):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.25
    font(p.add_run(text), size, color, bold, italic); return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(.375); p.paragraph_format.first_line_indent = Inches(-.188)
    p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.25
    font(p.add_run(text)); return p

def callout(doc, title, text, accent=GOLD):
    t = doc.add_table(rows=1, cols=1); geometry(t, [9360]); shade(t.cell(0,0), CREAM)
    p = t.cell(0,0).paragraphs[0]
    font(p.add_run(title + "\n"), 10, accent, True); font(p.add_run(text), 9.7)
    return t

def setup(title, subtitle, status):
    doc = Document(); sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.right_margin = sec.bottom_margin = sec.left_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(.492)
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
    for name, size, before, after in (("Heading 1",16,18,10),("Heading 2",13,14,7),("Heading 3",12,10,5)):
        style = doc.styles[name]; style.font.name = "Calibri"; style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(GOLD if name == "Heading 1" else NAVY)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
    hp = sec.header.paragraphs[0]; font(hp.add_run("AUCT.IO  |  ENTREGA 3"), 8.5, MUTED, True)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(fp.add_run("Grupo 15  ·  "), 8.5, MUTED)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); fp._p.append(fld)

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    font(p.add_run("INFORME DE ENTREGA"), 10, GOLD, True)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    font(p.add_run(title), 25, NAVY, True)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(14)
    font(p.add_run(subtitle), 13.5, MUTED)
    table(doc, ["Fecha", "Entorno", "Estado"], [["28/06/2026", "Android · Render · Azure SQL", status]], [1800, 3900, 3660], 9)
    return doc

def endpoint_inventory():
    source = (ROOT / "backend" / "index.js").read_text(encoding="utf-8")
    routes = re.findall(r'app\.(get|post|patch|put|delete)\("([^\"]+)"', source)
    groups = {"Autenticación y usuarios": [], "Subastas y pujas": [], "Administración": [], "Pagos, compras y multas": [], "Consignación y avisos": [], "Sistema": []}
    for method, path in routes:
        item = f"{method.upper()} {path}"
        if path.startswith("/api/auth") or path.startswith("/api/users"):
            groups["Autenticación y usuarios"].append(item)
        elif "/admin/" in path:
            groups["Administración"].append(item)
        elif any(key in path for key in ("payment-methods", "purchases", "fines")):
            groups["Pagos, compras y multas"].append(item)
        elif any(key in path for key in ("products", "notifications")):
            groups["Consignación y avisos"].append(item)
        elif any(key in path for key in ("auctions", "bids", "catalog-items", "subastas")):
            groups["Subastas y pujas"].append(item)
        else:
            groups["Sistema"].append(item)
    return routes, groups

routes, groups = endpoint_inventory()

tech = setup("Informe técnico Auct.io", "Arquitectura, componentes, datos, APIs y criterios de defensa", "Verificado para demo")
callout(tech, "RESUMEN EJECUTIVO", "Aplicación Android nativa conectada a una API Node.js/Express en Render y Azure SQL. Las pujas se confirman en transacción y se propagan por WebSocket. El cierre de cada lote es atómico y el último lote finaliza la subasta y libera a sus participantes.")
heading(tech, "1. Arquitectura de la solución", 1)
table(tech, ["Capa", "Tecnología", "Responsabilidad"], [
    ["Frontend", "Android · Java 11 · XML · SDK 36", "Interfaz móvil, validación previa, navegación, estados y consumo HTTP/WebSocket."],
    ["Backend", "Node.js · Express 5 · mssql · ws", "Reglas de negocio, autorizaciones, transacciones, temporizadores y eventos en tiempo real."],
    ["Base de datos", "Azure SQL Server", "Persistencia relacional de usuarios, subastas, lotes, pujas, ventas, pagos, multas y avisos."],
    ["Despliegue", "Render Web Service", "API pública HTTPS/WSS: https://auct-io-api.onrender.com."],
], [1500, 2500, 5360])
paragraph(tech, "Flujo de comunicación: Android → REST/JSON para operaciones; Android ↔ WebSocket para actualizaciones del catálogo; Backend → Azure SQL para lecturas y escrituras transaccionales.")

heading(tech, "2. Frontend Android", 1, page=True)
for text in [
    "Application ID com.example.clase4; minSdk 29, targetSdk/compileSdk 36 y versión 1.0.",
    "Pantallas principales: acceso y registro, inicio, subastas, catálogo en vivo, puja, pagos, compras, multas, consignación, avisos, estadísticas, perfil y administración.",
    "Diseño XML con paleta azul petróleo, dorado y crema; diálogos FeedbackDialog consistentes; launcher adaptativo con martillo de subastas.",
    "ApiConfig inyecta la URL de Render desde BuildConfig; HttpURLConnection y WebSocket cubren el transporte productivo.",
    "El panel administrativo presenta KPIs, badges y colas accionables; el ABM no exige copiar identificadores técnicos.",
]: bullet(tech, text)
heading(tech, "Tratamiento horario", 2)
paragraph(tech, "Los formularios y textos operativos usan America/Argentina/Buenos_Aires. Las fechas de subasta se interpretan como hora civil argentina GMT-3; los contadores se calculan en servidor y no dependen del reloj del teléfono.")

heading(tech, "3. Backend y reglas de negocio", 1, page=True)
table(tech, ["Regla", "Implementación"], [
    ["Categorías", "común < especial < plata < oro < platino; el cliente debe igualar o superar la categoría."],
    ["Puja", "Debe superar la mejor oferta; mínimo +1% y máximo +20% del precio base, excepto oro/platino."],
    ["Exclusividad", "Una sesión de subasta por usuario; puede salir si aún no ofertó."],
    ["Garantía", "Medio verificado y compatible con moneda; control de cheque certificado y obligaciones vencidas."],
    ["Cierre", "Transacción SERIALIZABLE: ganador/empresa, venta, nuevo dueño y lote se confirman juntos."],
    ["Sin pujas", "La empresa (cliente 9000007) compra al precio base, conforme a la consigna."],
    ["Tiempo real", "WebSocket /ws difunde catálogo, mejor oferta, historial y reloj por lote."],
], [1800, 7560])
heading(tech, "Temporizador y reconciliación", 2)
paragraph(tech, "Un proceso cada cinco segundos inicia subastas vencidas, cierra lotes cuyo reloj llegó a cero y corrige subastas con todos sus lotes finalizados. El proceso funciona aunque no haya una pantalla abierta. Cada puja válida reinicia únicamente el reloj de su lote.")

heading(tech, "4. Modelo de datos", 1, page=True)
table(tech, ["Entidad", "Relación principal"], [
    ["Users / Clients / Owners / Employees", "Identidad base y especializaciones por rol."],
    ["Auctions / Catalogs / CatalogItems", "Evento, catálogo y lotes con precio base/comisión."],
    ["Attendees / Bids", "Participación y secuencia completa de ofertas."],
    ["AuctionRecords", "Venta adjudicada, comprador, dueño anterior, importe, pago, envío y retiro."],
    ["PaymentMethods / Fines", "Garantías, monedas, verificaciones e incumplimientos."],
    ["Products / Photos / Insurances", "Consignación, seis o más imágenes, depósito y cobertura."],
    ["Notifications", "Avisos persistentes de acciones relevantes."],
], [2600, 6760])
callout(tech, "ZONA HORARIA", "Los defaults de eventos relevantes en Azure SQL usan DATEADD(HOUR,-3,SYSUTCDATETIME()). La auditoría confirmó cero defaults operativos fuera de GMT-3.", GREEN)

heading(tech, "5. API REST y WebSocket", 1, page=True)
paragraph(tech, f"El backend expone {len(routes)} rutas REST declaradas en Express, además de /ws. Los códigos más usados son 200/201/202 para éxito, 400 para validación, 401/403 para identidad o regla de acceso, 404 para inexistencia y 409 para conflictos de estado.")
table(tech, ["Endpoint clave", "Uso"], [
    ["POST /api/auth/login", "Autentica cliente o empleado y devuelve sesión."],
    ["GET /api/clients/:clientId/auctions", "Lista subastas y explica si puede pujar."],
    ["GET /api/auctions/:auctionId/catalog", "Catálogo completo, mejor oferta, mi oferta y reloj por lote."],
    ["POST /api/bids", "Valida y confirma una puja de forma serializada."],
    ["POST /api/clients/:clientId/active-auction/release", "Libera una conexión cuando la regla lo permite."],
    ["POST /api/admin/auctions", "Crea subasta y catálogo en una transacción."],
    ["POST /api/admin/auctions/:auctionId/items/:itemId/close", "Cierre manual atómico del lote."],
    ["GET /api/admin/dashboard", "KPIs del panel interno."],
    ["WS /ws?auctionId=&clientId=", "Actualizaciones simultáneas del catálogo."],
], [4400, 4960])

for index, (group, items) in enumerate(groups.items(), start=6):
    heading(tech, f"{index}. Inventario: {group}", 1, page=True)
    for item in items: bullet(tech, item)

heading(tech, "12. Seguridad, consistencia y límites", 1, page=True)
for text in [
    "Las operaciones administrativas exigen token de empleado mediante requireEmployee.",
    "Las pujas se serializan en SQL y el cliente deshabilita el envío hasta recibir confirmación.",
    "El cierre usa bloqueo UPDLOCK/HOLDLOCK y aislamiento SERIALIZABLE para evitar doble adjudicación.",
    "Las credenciales SQL se reciben por variables de entorno y no se documentan ni incluyen en el APK.",
    "La autenticación actual usa un token demostrativo codificado, no un JWT firmado. En producción debe reemplazarse por OAuth/JWT con expiración y rotación.",
    "Pagos, aseguradora, correo, streaming y transferencias se representan mediante estados y avisos; conectar proveedores reales queda fuera del alcance académico.",
]: bullet(tech, text)

heading(tech, "13. Evidencia de pruebas", 1)
table(tech, ["Prueba", "Resultado"], [
    ["Smoke REST", "Rutas públicas, cliente y administrador: aprobado."],
    ["ABM de subastas", "Alta, modificación, listado y baja segura: aprobado."],
    ["Cierre integral", "1 lote con puja + 1 sin puja; 2 ventas; subasta cerrada; sesión liberada."],
    ["Concurrencia", "2 clientes reciben el mismo catálogo por WebSocket."],
    ["Base productiva", "0 cierres inconsistentes, 0 lotes vendidos sin venta, defaults GMT-3."],
    ["Android", "assembleDebug, unit tests y lintDebug: aprobados."],
], [3000, 6360])

heading(tech, "14. Preguntas probables de defensa", 1, page=True)
qa = [
    ("¿Dónde está la lógica de negocio?", "En Express; Android valida para UX, pero el backend vuelve a validar y Azure SQL asegura la transacción."),
    ("¿Cómo evitan dos ganadores?", "La puja y el cierre usan operaciones serializadas y locks en SQL; solo el mejor registro se marca ganador."),
    ("¿Qué pasa sin ofertas?", "Al vencer el lote se crea una venta pagada a nombre del cliente empresa por el precio base."),
    ("¿Cómo se libera un usuario?", "Al finalizar el último lote se cierra Auctions y se eliminan todas las sesiones en memoria de esa subasta."),
    ("¿Por qué WebSocket?", "Evita polling continuo y difunde cambios de oferta, reloj e historial a todos los participantes."),
    ("¿Qué escalaría para producción?", "JWT/OAuth, Redis para sesiones y pub/sub multi-instancia, almacenamiento Blob para fotos y observabilidad centralizada."),
]
table(tech, ["Pregunta", "Respuesta sugerida"], qa, [3000, 6360], 8.8)
tech.save(TECH_OUT)

summary = setup("Resumen de cambios y defensa", "Cómo funcionan las subastas después de la corrección", "Listo para el 29/06")
callout(summary, "RESULTADO", "El usuario ya no queda cautivo cuando terminó el catálogo. La subasta se cierra al finalizar su último lote, libera sesiones y conserva ventas, avisos y trazabilidad.", GREEN)
heading(summary, "1. Cambio funcional aplicado", 1)
for text in [
    "Cada lote tiene reloj independiente y una puja reinicia solo ese reloj.",
    "Al vencer, el lote se adjudica al mejor postor dentro de una transacción.",
    "Si no hubo pujas, Auct.io Empresa compra al precio base, como exige la consigna.",
    "Cuando no quedan lotes abiertos, Auctions pasa a cerrada y se liberan todos los usuarios vinculados, hayan ganado o no.",
    "Un reconciliador ejecuta el flujo aunque ningún usuario tenga la app abierta.",
    "Todas las fechas operativas y la agenda se interpretan en Argentina GMT-3.",
]: bullet(summary, text)
heading(summary, "2. Subastas preparadas", 1)
table(summary, ["ID", "Inicio Argentina", "Duración/lote", "Categoría · moneda", "Contenido"], [
    ["17", "29/06 · 18:00", "120 min", "común · pesos", "Arte y colección · 2 lotes"],
    ["18", "29/06 · 19:30", "120 min", "plata · pesos", "Diseño y joyas · 2 lotes"],
    ["19", "29/06 · 21:00", "90 min", "oro · dólares", "Colección premium · 2 lotes"],
], [650, 1900, 1500, 2200, 3110])
paragraph(summary, "Cada producto incluye seis fotografías. Los eventos se solapan para que haya una alternativa activa durante la franja de defensa entre las 18:00 y las 22:00.")
heading(summary, "3. Flujo breve para mostrar", 1, page=True)
steps = [
    "Ingresar como administrador y mostrar KPIs, pendientes y Gestionar subastas.",
    "Abrir una subasta de defensa y comprobar fecha/hora GMT-3 y sus dos lotes.",
    "Ingresar con dos clientes en dispositivos o sesiones diferentes.",
    "Pujar en un lote y observar actualización WebSocket, historial y reinicio del reloj.",
    "Dejar el segundo lote sin ofertas o cerrarlo desde administración.",
    "Finalizar el lote ofertado y comprobar ganador, compra pendiente y aviso.",
    "Confirmar que el lote sin pujas fue comprado por la empresa al precio base.",
    "Volver al listado e ingresar a otra subasta: la sesión anterior ya fue liberada.",
]
for idx, text in enumerate(steps, 1): paragraph(summary, f"{idx}. {text}")
heading(summary, "4. Controles realizados", 1)
for text in [
    "Cierre integral con y sin pujas: aprobado.", "Liberación de ganador: aprobada.",
    "ABM administrativo: aprobado.", "WebSocket con dos clientes: aprobado.",
    "Consistencia Azure SQL y GMT-3: aprobada.", "Compilación, unit tests y lint: aprobados.",
    "Panel administrativo y formulario revisados en emulador; corregido el formato horario de SQL Server.",
]: bullet(summary, text)
callout(summary, "RECOMENDACIÓN DE DEMO", "Usar primero la subasta #17. Las #18 y #19 quedan como respaldo o para demostrar categorías y moneda en dólares.")
summary.save(SUMMARY_OUT)
print(TECH_OUT)
print(SUMMARY_OUT)
